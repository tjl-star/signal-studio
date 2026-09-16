import csv
import json
import os
import urllib.parse
import urllib.request
from datetime import date, timedelta

BASE = "http://101.132.68.174:8123/skill/"
START = "2026-07-05"
END = "2026-08-04"
OUT = os.path.join(os.environ.get("USERPROFILE", r"C:\Users\tjldq"), "Desktop", "首页流量与转化漏斗")


def fetch(endpoint, params):
    url = BASE + endpoint + "?" + urllib.parse.urlencode(params)
    with urllib.request.urlopen(url, timeout=120) as response:
        return json.load(response)


def write_csv(path, rows):
    keys = sorted({key for row in rows for key in row}) if rows else []
    with open(path, "w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=keys, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def main():
    os.makedirs(OUT, exist_ok=True)
    params = {"startDate": START, "endDate": END}
    datasets = {
        "drama_conversion_daily": fetch("dramaConversion", params),
        "custom_navigation_daily": fetch("customNavigation", params),
        "custom_navigation_detail": fetch("customNavigationDetail", params),
        "available_play_daily": fetch("availablePlayData", params),
        "search_click_conversion_daily": fetch("searchClickConversion", params),
    }

    for name, rows in datasets.items():
        write_csv(os.path.join(OUT, name + ".csv"), rows)

    metadata = {
        "data_range": {"start": START, "end": END, "inclusive": True},
        "scope": "all users / all available channels and pages",
        "sources": {
            "drama_conversion_daily": "/skill/dramaConversion",
            "custom_navigation_daily": "/skill/customNavigation",
            "custom_navigation_detail": "/skill/customNavigationDetail",
            "available_play_daily": "/skill/availablePlayData",
            "search_click_conversion_daily": "/skill/searchClickConversion",
        },
        "notes": [
            "customNavigation and customNavigationDetail were exported without page_name filtering because the API returned no rows for page_name=首页; the files retain the original page_name/page_id fields.",
            "The datasets are kept at their original grains and are not joined across endpoints.",
        ],
        "row_counts": {name: len(rows) for name, rows in datasets.items()},
    }
    with open(os.path.join(OUT, "metadata.json"), "w", encoding="utf-8") as handle:
        json.dump(metadata, handle, ensure_ascii=False, indent=2)

    try:
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter

        workbook = Workbook()
        workbook.remove(workbook.active)
        for name, rows in datasets.items():
            sheet = workbook.create_sheet(name[:31])
            keys = sorted({key for row in rows for key in row}) if rows else []
            sheet.append(keys)
            for row in rows:
                sheet.append([row.get(key) for key in keys])
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for index, key in enumerate(keys, 1):
                width = min(max(len(str(key)) + 2, 12), 28)
                sheet.column_dimensions[get_column_letter(index)].width = width
        workbook.save(os.path.join(OUT, "首页流量与转化漏斗_20260705-20260804.xlsx"))
    except ImportError:
        pass


if __name__ == "__main__":
    main()
