from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "data-provider缺失字段流程图.docx"
PNG = ROOT / "data-provider缺失字段流程图.png"

groups = [
    ("播放数据 → 播放排行榜", ["播放人数", "榜单类型"]),
    ("功能数据 → 搜索 → 搜索整体数据", ["搜索人数", "内容类型", "题材", "用户类型"]),
    ("功能数据 → 搜索 → 搜索漏斗", ["搜索用户数", "详情用户数", "首帧播放用户数", "五分钟有效播放用户数", "完成播放用户数", "搜索到详情转化率", "详情到首帧播放转化率", "有效播放率", "完成率"]),
    ("页面数据 → 频道页流量漏斗", ["首页用户数", "频道页用户数", "详情页用户数", "播放用户数", "频道转化率", "昨日环比"]),
    ("页面数据 → 首页板块数据", ["板块播放次数", "板块有效播放次数", "板块点击转化率"]),
    ("页面数据 → 推荐组件数据 → 横幅点击数据", ["横幅播放次数", "横幅有效播放次数"]),
]

font_candidates = [
    Path("C:/Windows/Fonts/msyh.ttc"),
    Path("C:/Windows/Fonts/simhei.ttf"),
    Path("C:/Windows/Fonts/simsun.ttc"),
]
font_path = next((p for p in font_candidates if p.exists()), None)
if font_path is None:
    raise RuntimeError("找不到中文字体")

def font(size):
    return ImageFont.truetype(str(font_path), size=size)

def text_box(draw, xy, text, fnt, fill, outline, radius=14, padding=18):
    x, y, w, h = xy
    draw.rounded_rectangle((x, y, x + w, y + h), radius=radius, fill=fill, outline=outline, width=2)
    bbox = draw.multiline_textbbox((0, 0), text, font=fnt, spacing=6, align="center")
    tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
    draw.multiline_text((x + (w - tw) / 2, y + (h - th) / 2 - 2), text, font=fnt, fill="#17324d", spacing=6, align="center")

def make_flowchart():
    W = 1800
    group_y = [120, 315, 510, 705, 900, 1095]
    img = Image.new("RGB", (W, 1270), "#ffffff")
    draw = ImageDraw.Draw(img)
    title_font = font(40)
    group_font = font(25)
    field_font = font(24)
    draw.text((40, 28), "数据接口缺失字段流程图", font=title_font, fill="#17324d")
    draw.text((40, 82), "Quick BI 页面位置 → 需要新增的中文接口字段", font=font(22), fill="#657789")
    text_box(draw, (40, 690, 280, 90), "需要新增\n接口字段", font(28), "#e8f1f8", "#76a8c9")
    for idx, (path, fields) in enumerate(groups):
        y = group_y[idx]
        text_box(draw, (430, y, 360, 92), path, group_font, "#f1f5f8", "#a5b5c2")
        draw.line((320, 735, 430, y + 46), fill="#b7c4ce", width=3)
        field_w = 300
        field_x = [840, 1160, 1480]
        for j, field in enumerate(fields):
            row = j // 3
            col = j % 3
            fy = y + 2 + row * 62
            text_box(draw, (field_x[col], fy, field_w, 50), field, field_font, "#ffffff", "#93b5c9", radius=12)
            draw.line((790, y + 46, field_x[col], fy + 25), fill="#c1ccd4", width=2)
    img.save(PNG)

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")

def set_table_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.first_child_found_in("w:tcW")
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(int(width * 1440)))
            tcW.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

make_flowchart()
doc = Document()
section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width = Inches(11.69)
section.page_height = Inches(8.27)
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.1

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(4)
r = title.add_run("data-provider 缺失字段需求清单")
r.font.name = "Microsoft YaHei"
r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
r.font.size = Pt(24)
r.font.bold = True
r.font.color.rgb = RGBColor(23, 50, 77)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.paragraph_format.space_after = Pt(12)
sr = sub.add_run("Quick BI 页面位置与新增中文接口字段对应关系")
sr.font.name = "Microsoft YaHei"
sr._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
sr.font.size = Pt(11)
sr.font.color.rgb = RGBColor(101, 119, 137)

doc.add_picture(str(PNG), width=Inches(9.9))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("字段对应表", level=1)
table = doc.add_table(rows=1, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
set_table_widths(table, [3.0, 7.0])
headers = ["需要新增的接口字段", "Quick BI 位置"]
for cell, text in zip(table.rows[0].cells, headers):
    cell.text = text
    set_cell_shading(cell, "DCEAF3")
    for p in cell.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.bold = True
            run.font.name = "Microsoft YaHei"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            run.font.size = Pt(10)
for path, fields in groups:
    for field in fields:
        cells = table.add_row().cells
        cells[0].text = field
        cells[1].text = path
        for cell in cells:
            set_cell_margins(cell)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = "Microsoft YaHei"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
                    run.font.size = Pt(9.5)

doc.add_paragraph()
note = doc.add_paragraph()
note.paragraph_format.space_before = Pt(4)
nr = note.add_run("字段数量：26 个。文档仅包含当前缺失字段，不包含已有字段。")
nr.font.name = "Microsoft YaHei"
nr._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
nr.font.size = Pt(9.5)
nr.font.color.rgb = RGBColor(101, 119, 137)

doc.save(OUT)
print(OUT)
